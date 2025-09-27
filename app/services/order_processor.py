# app/services/order_processor.py
from __future__ import annotations
import os, re, tempfile
from typing import Optional, List, Dict, Any

import pandas as pd
import openpyxl  # đọc hyperlink ẩn
from pathlib import Path

# ============================================================
# Paths (ổn định & có thể cấu hình)
# ============================================================
PROJECT_ROOT  = Path(__file__).resolve().parents[2]          # .../ps_tool
TEMPLATE_ROOT = Path(os.environ.get("TEMPLATE_ROOT", PROJECT_ROOT / "storage" / "templates")).resolve()
OUTPUTS_DIR   = (PROJECT_ROOT / "storage" / "outputs").resolve()
TMP_DIR       = (PROJECT_ROOT / "storage" / "tmp").resolve()
DEFAULT_EXCEL = (PROJECT_ROOT / "data.xlsx").resolve()

# Nếu module tải Drive chưa sẵn có, giữ API nhưng không làm vỡ luồng.
try:
    from app.services.google_drive import download_from_share_link
except Exception:
    def download_from_share_link(url: str) -> bytes:
        raise RuntimeError("Google Drive downloader chưa cấu hình.")

# ============================================================
# Utils chung
# ============================================================
def _ensure_dir(p: str | Path): Path(p).mkdir(parents=True, exist_ok=True)

def vn_key(s: str) -> str:
    """Lower + bỏ dấu để so khớp tiêu đề cột không phân biệt dấu."""
    if s is None: return ""
    s = str(s)
    try:
        import unicodedata as ud
        s = ud.normalize("NFD", s)
        s = "".join(ch for ch in s if ud.category(ch) != "Mn")
    except Exception:
        pass
    return s.lower().strip()

def looks_like_url(x: Any) -> bool:
    if not x: return False
    s = str(x).strip()
    return s.startswith("http://") or s.startswith("https://")

def extract_urls(cell: Any) -> List[str]:
    if cell is None: return []
    s = str(cell)
    return re.findall(r'(https?://\S+)', s)

def prefer_gdrive(urls: List[str]) -> List[str]:
    g = [u for u in urls if "drive.google.com" in u]
    return g if g else urls

def clean_order_id(val: Any) -> str:
    """Loại '.0' khi Excel đọc số; giữ format gốc."""
    if val is None: return ""
    s = str(val).strip()
    if re.fullmatch(r"\d+\.0", s): s = s[:-2]
    return s

INVALID_FS_CHARS = r'[<>:"/\\|?*\n\r]'
def _sanitize_filename(name: str, max_len: int = 220) -> str:
    s = str(name)
    s = re.sub(INVALID_FS_CHARS, "-", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s[:max_len]

def _norm_cols(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    return df

def _col_by_letter(df: pd.DataFrame, letter: str) -> Optional[str]:
    idx = ord(letter.upper()) - ord("A")
    return df.columns[idx] if 0 <= idx < len(df.columns) else None

def _pick_by_letter_or_alias(df: pd.DataFrame, prefer_letter: str, aliases: List[str]) -> Optional[str]:
    col = _col_by_letter(df, prefer_letter)
    if col is not None:
        return col
    norm = {vn_key(c): c for c in df.columns}
    for a in aliases:
        c = norm.get(vn_key(a))
        if c: return c
    return None

def _right_without_first_char(s: str) -> str:
    s = "" if s is None else str(s)
    return s[1:] if len(s) > 0 else ""

def _compute_counts_for_O(df: pd.DataFrame, o_col: str, count_range_max: int = 1000):
    """COUNTIF giống Excel: tổng và thứ tự xuất hiện theo cột O."""
    work = df.copy()
    o_series = work[o_col].astype(str)
    upto = min(len(work), max(1, count_range_max - 1))  # Excel O2:O1000 ~ pandas 0..upto-1
    totals = o_series.iloc[:upto].value_counts()
    work["_o_key"]   = o_series
    work["_o_total"] = work["_o_key"].map(totals).fillna(0).astype(int)
    work["_o_cum"]   = work.groupby("_o_key").cumcount() + 1
    return work[["_o_total", "_o_cum"]]

# ---------- đọc hyperlink ẩn ----------
def _build_ws_header_map(ws) -> Dict[str, int]:
    """
    Map header text -> column index (1-based).
    Không phụ thuộc thuộc tính col_idx (incompatible với ReadOnlyCell).
    """
    header_rows = ws.iter_rows(min_row=1, max_row=1, values_only=False)
    try:
        header_cells = next(header_rows)
    except StopIteration:
        return {}
    mapping: Dict[str, int] = {}
    for i, c in enumerate(header_cells, start=1):
        val = c.value
        if val is None:
            continue
        key = str(val).strip()
        if key:
            mapping[key] = i  # 1-based index
    return mapping

def _hyperlinks_in_row(ws, row_idx: int, col_indexes: Optional[List[int]] = None) -> List[str]:
    """
    Lấy hyperlink.target ở dòng Excel (1-based).
    - col_indexes: nếu truyền → chỉ quét các cột đó; nếu None → quét toàn dòng.
    """
    urls = []
    if ws is None or row_idx < 2:
        return urls
    cells = [ws.cell(row=row_idx, column=ci) for ci in col_indexes] if col_indexes else ws[row_idx]
    for cell in cells:
        try:
            h = cell.hyperlink
            if h and h.target:
                urls.append(str(h.target))
        except Exception:
            pass
    return urls

# ============================================================
# Tìm template sda.psd "thông minh"
# ============================================================
def _find_dir_case_insensitive(base: Path, name: str) -> Optional[Path]:
    """Tìm thư mục con khớp tên (case-insensitive)."""
    if not base.is_dir(): return None
    target = str(name).strip().lower()
    for d in base.iterdir():
        if d.is_dir() and d.name.lower() == target:
            return d
    return None

def _normalize_size_variants(size_text: str) -> List[str]:
    """
    Biến thể size:
      - giữ nguyên (ví dụ "Adult M")
      - viết hoa
      - bỏ tiền tố nhóm (Adult/Youth/Kid/Infant) => "M"
      - bản viết hoa của biến thể đơn giản
    """
    s = (size_text or "").strip()
    variants = []
    if not s:
        return variants
    variants.append(s)
    variants.append(s.upper())
    s_simple = re.sub(r"^(Adult|Youth|Kid|Infant)\s+", "", s, flags=re.I).strip()
    if s_simple and s_simple not in variants:
        variants.append(s_simple)
    s_simple_u = s_simple.upper()
    if s_simple_u and s_simple_u not in variants:
        variants.append(s_simple_u)
    return variants

def _resolve_template_path(type_name: str, size_text: str) -> Optional[str]:
    """
    Tìm .../<type>/<size>/sda.psd:
      - so khớp thư mục không phân biệt hoa/thường
      - thử nhiều biến thể size (Adult M -> M, v.v.)
      - fallback: .../sda.psd trong type, rồi ROOT, rồi PROJECT_ROOT
    """
    troot = TEMPLATE_ROOT
    type_dir = _find_dir_case_insensitive(troot, str(type_name).strip()) or (troot / str(type_name).strip())
    size_variants = _normalize_size_variants(size_text)

    # 1) Thử trong thư mục type (nếu có)
    if type_dir.is_dir():
        for s in size_variants:
            sdir = _find_dir_case_insensitive(type_dir, s) or (type_dir / s)
            psd = sdir / "sda.psd"
            if psd.is_file():
                return str(psd)
        # Fallback: .../<type>/sda.psd
        psd = type_dir / "sda.psd"
        if psd.is_file():
            return str(psd)

    # 2) ROOT/sda.psd
    psd = troot / "sda.psd"
    if psd.is_file():
        return str(psd)
    # 3) PROJECT_ROOT/sda.psd
    psd = PROJECT_ROOT / "sda.psd"
    if psd.is_file():
        return str(psd)
    return None

# ============================================================
# Core: đọc Excel
# ============================================================
def _pick_columns(df: pd.DataFrame) -> Dict[str, Optional[str]]:
    """Chọn cột theo chữ cái ưu tiên + alias đa dạng (có/không dấu)."""
    J = _pick_by_letter_or_alias(df, "J", ["Loại áo","Loai ao","Type","Loại","Loai"])
    N = _pick_by_letter_or_alias(df, "N", ["Màu","Mau","Color"])
    M = _pick_by_letter_or_alias(df, "M", ["Size","Kích cỡ","Kich co","Cỡ","Co"])
    O = _pick_by_letter_or_alias(df, "O", ["Order ID","Mã đơn","Ma don","Mã đơn hàng","Ma don hang","Order","Order Number"])
    # link có thể nằm ở L/K hoặc bất kỳ cột nào có URL
    link = _pick_by_letter_or_alias(df, "L", ["File thiết kế","File thiet ke","Link file thiết kế","Link"])
    if link is None:
        # fallback: tìm cột có nhiều URL nhất
        url_counts = []
        for c in df.columns:
            cnt = 0
            for v in df[c].head(200):
                if looks_like_url(v) or extract_urls(v):
                    cnt += 1
            url_counts.append((c, cnt))
        url_counts.sort(key=lambda x: x[1], reverse=True)
        if url_counts and url_counts[0][1] > 0:
            link = url_counts[0][0]
    return {"J": J, "N": N, "M": M, "O": O, "L": link}

def _choose_sheet(xl: pd.ExcelFile, prefer: Optional[str]) -> str:
    sheets = xl.sheet_names
    if prefer and prefer in sheets:
        return prefer
    best, best_score = sheets[0], -1
    for sh in sheets:
        try:
            df = xl.parse(sh, nrows=30)
        except Exception:
            continue
        df = _norm_cols(df)
        picks = _pick_columns(df)
        score = sum(1 for k in ("J","N","M","O") if picks.get(k)) + (1 if picks.get("L") else 0)
        if score > best_score:
            best, best_score = sh, score
    return best

# ---------- helpers bắt link nâng cao ----------
def _extract_urls_from_row_text(row_like) -> List[str]:
    """Lấy mọi URL xuất hiện trong TEXT của toàn bộ ô trên 1 dòng."""
    urls: List[str] = []
    # chấp nhận cả dict, Series, list/tuple
    if hasattr(row_like, "to_dict"):
        it = row_like.to_dict().values()
    elif isinstance(row_like, dict):
        it = row_like.values()
    else:
        it = list(row_like)
    for v in it:
        urls.extend(extract_urls(v))
    # unique, giữ thứ tự
    seen, out = set(), []
    for u in urls:
        if u not in seen:
            seen.add(u); out.append(u)
    return out

def _extract_drive_ids_from_row(row_like) -> List[str]:
    """Bắt các Google Drive ID rời rạc/ẩn trong cả dòng."""
    if hasattr(row_like, "to_dict"):
        it = row_like.to_dict().values()
    elif isinstance(row_like, dict):
        it = row_like.values()
    else:
        it = list(row_like)

    ids: List[str] = []
    rx = [
        r"/d/([A-Za-z0-9_-]{20,})",           # .../file/d/<ID>/
        r"[?&]id=([A-Za-z0-9_-]{20,})",       # ...?id=<ID>
        r"\b([A-Za-z0-9_-]{20,})\b",          # ID rời rạc (>=20)
    ]
    for v in it:
        s = "" if v is None else str(v)
        for pat in rx:
            for m in re.findall(pat, s):
                if len(m) >= 20:
                    ids.append(m)
    # unique, giữ thứ tự
    seen, out = set(), []
    for i in ids:
        if i not in seen:
            seen.add(i); out.append(i)
    return out

# ---------- main ----------
def process_order_from_excel(excel_path: Optional[str] = None,
                             sheet: Optional[str] = None,
                             limit: Optional[int] = None,
                             count_range_max: int = 1000) -> List[Dict[str, Any]]:
    """
    Trả về mảng item với shape cũ UI cần:
      order_id, template, output_path, status, error
    + các trường nội bộ _psd_path, _img_path để psd_handler dùng.
    """
    excel = Path(excel_path).resolve() if excel_path else DEFAULT_EXCEL
    if not excel.is_file():
        raise FileNotFoundError(f"Không thấy Excel: {excel}")

    xl = pd.ExcelFile(excel, engine="openpyxl")
    target_sheet = _choose_sheet(xl, sheet)
    df = xl.parse(target_sheet)
    df = _norm_cols(df)

    # Nạp worksheet để đọc hyperlink ẩn (mở thường để có đủ thuộc tính)
    wb = openpyxl.load_workbook(excel, data_only=True, read_only=False)
    ws = wb[target_sheet]
    ws_header_map = _build_ws_header_map(ws)

    picks = _pick_columns(df)
    miss = [k for k in ("J","N","M","O") if not picks.get(k)]
    if miss:
        raise KeyError(f"Không xác định được cột: {', '.join(miss)} (J/N/M/O).")

    counts = _compute_counts_for_O(df, picks["O"], count_range_max=count_range_max)
    _ensure_dir(OUTPUTS_DIR); _ensure_dir(TMP_DIR)

    data_rows = df if not limit or limit <= 0 else df.head(limit)
    items: List[Dict[str, Any]] = []

    for idx, row in data_rows.iterrows():
        try:
            jv = str(row[picks["J"]]).strip()     # Loại áo
            nv = str(row[picks["N"]]).strip()     # Màu
            mv = str(row[picks["M"]]).strip()     # Size
            ov = clean_order_id(row.get(picks["O"]))  # Order ID
            template = f"{jv}/{mv}"

            # Tên file output theo công thức Excel
            display_name = f"{jv}-{nv}-{mv}-{_right_without_first_char(str(ov))}-{int(counts.loc[idx, '_o_cum'])}/{int(counts.loc[idx, '_o_total'])}"
            safe_name    = _sanitize_filename(display_name)
            out_path     = str(OUTPUTS_DIR / f"{safe_name}.jpg")

            # Tìm sda.psd (linh hoạt)
            psd_path = _resolve_template_path(jv, mv)
            if not psd_path:
                items.append({
                    "order_id": ov, "template": template, "output_path": None,
                    "status": "ERROR", "error": f"Không tìm thấy sda.psd cho '{template}'"
                })
                continue

            # ===== Lấy link file thiết kế (mạnh tay) =====
            # 1) Từ text ô cột link (nếu pick được)
            links = prefer_gdrive(extract_urls(row.get(picks["L"]))) if picks.get("L") else []

            # 2) Nếu chưa có, quét toàn bộ text của dòng
            if not links:
                links = prefer_gdrive(_extract_urls_from_row_text(row))

            # 3) Nếu vẫn chưa có, lấy hyperlink ẩn (ưu tiên cột link nếu biết; không thì cả dòng)
            if not links:
                excel_row_idx = int(idx) + 2  # header ở dòng 1
                link_col_indexes = []
                if picks.get("L") and picks["L"] in ws_header_map:
                    link_col_indexes = [ws_header_map[picks["L"]]]
                hlinks = _hyperlinks_in_row(ws, excel_row_idx, link_col_indexes or None)
                if hlinks:
                    links = prefer_gdrive(hlinks)

            # 4) Nếu vẫn chưa có, dựng URL từ Google Drive ID rời rạc
            if not links:
                ids = _extract_drive_ids_from_row(row)
                if ids:
                    links = [f"https://drive.google.com/file/d/{ids[0]}/view"]

            design_link = links[0] if links else None
            if not design_link:
                items.append({
                    "order_id": ov, "template": template, "output_path": None,
                    "status": "SKIP", "error": "Thiếu link 'File thiết kế'"
                })
                continue

            # Tải file thiết kế -> ảnh tạm
            try:
                data = download_from_share_link(design_link)
            except Exception as e:
                items.append({
                    "order_id": ov, "template": template, "output_path": None,
                    "status": "SKIP", "error": f"Tải file thiết kế lỗi: {e}"
                })
                continue

            # Đoán đuôi ảnh
            try:
                import imghdr
                ext = imghdr.what(None, h=data) or "jpg"
            except Exception:
                ext = "jpg"
            fd, tmp_img = tempfile.mkstemp(prefix="design_", suffix=f".{ext}", dir=str(TMP_DIR))
            os.close(fd)
            with open(tmp_img, "wb") as f:
                f.write(data)

            # READY cho handler
            items.append({
                "order_id": ov,
                "template": template,
                "output_path": out_path,
                "status": "READY",
                "error": None,
                "_psd_path": psd_path,
                "_img_path": tmp_img,
                "_tmp_img": True,
            })

        except Exception as e:
            items.append({
                "order_id": "", "template": "", "output_path": None,
                "status": "ERROR", "error": str(e),
            })

    return items
